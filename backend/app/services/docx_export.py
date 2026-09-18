"""Geração de documento DOCX profissional a partir da transcrição."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ..database import Reuniao


def _formatar_tempo(segundos: float) -> str:
    segundos = int(segundos)
    h, resto = divmod(segundos, 3600)
    m, s = divmod(resto, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def gerar_docx_transcricao(reuniao: Reuniao) -> BytesIO:
    documento = Document()

    titulo = documento.add_heading(reuniao.titulo, level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_heading("Dados da reunião", level=2)
    tabela = documento.add_table(rows=0, cols=2)
    tabela.style = "Light Grid Accent 1"

    def _linha(rotulo: str, valor: str) -> None:
        celulas = tabela.add_row().cells
        celulas[0].text = rotulo
        celulas[1].text = valor

    _linha("Arquivo original", reuniao.nome_arquivo_original)
    _linha("Data de processamento", reuniao.criado_em.strftime("%d/%m/%Y %H:%M"))
    if reuniao.duracao_segundos:
        _linha("Duração", _formatar_tempo(reuniao.duracao_segundos))
    _linha("Idioma", reuniao.idioma)

    participantes = sorted({s.falante for s in reuniao.segmentos if s.falante})
    if participantes:
        _linha("Participantes identificados", ", ".join(participantes))
    else:
        _linha(
            "Participantes",
            "Identificação individual de participantes (diarização) não "
            "disponível nesta versão.",
        )

    documento.add_paragraph()
    documento.add_heading("Transcrição", level=2)

    if not reuniao.segmentos:
        documento.add_paragraph("Nenhum segmento de transcrição disponível.")

    for segmento in reuniao.segmentos:
        p = documento.add_paragraph()
        marca = p.add_run(f"[{_formatar_tempo(segmento.inicio_segundos)}] ")
        marca.bold = True
        marca.font.size = Pt(10)

        if segmento.falante:
            falante = p.add_run(f"{segmento.falante}: ")
            falante.bold = True

        texto = segmento.texto
        if segmento.baixa_confianca:
            texto = f"{texto} [trecho com baixa confiança - revisar]"
        p.add_run(texto)

    rodape = documento.add_paragraph()
    rodape.add_run(
        "\nDocumento gerado automaticamente pelo Transcritor CTCE a partir de "
        "processamento local de áudio. Recomenda-se revisão humana antes do uso "
        "oficial deste conteúdo."
    ).italic = True

    buffer = BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer
