"""Geração de documento DOCX da Ata/Relatório da reunião."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..database import AtaReuniao, Encaminhamento, Reuniao


def _formatar_tempo(segundos: float) -> str:
    segundos = int(segundos)
    h, resto = divmod(segundos, 3600)
    m, s = divmod(resto, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def _adicionar_secao(documento: Document, titulo: str, conteudo: str | None, texto_vazio: str) -> None:
    documento.add_heading(titulo, level=2)
    documento.add_paragraph(conteudo.strip() if conteudo and conteudo.strip() else texto_vazio)


def gerar_docx_ata(reuniao: Reuniao, ata: AtaReuniao, encaminhamentos: list[Encaminhamento]) -> BytesIO:
    documento = Document()

    titulo = documento.add_heading(f"Ata da Reunião - {reuniao.titulo}", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if ata.gerada_por_ia:
        aviso = documento.add_paragraph()
        aviso.add_run(
            "Este documento contém trechos com rascunho gerado por modelo de linguagem "
            "local (processamento offline). Revise cuidadosamente a fidelidade ao "
            "conteúdo da reunião antes de qualquer uso oficial."
        ).italic = True

    documento.add_heading("Dados da reunião", level=2)
    tabela = documento.add_table(rows=0, cols=2)
    tabela.style = "Light Grid Accent 1"

    def _linha(rotulo: str, valor: str) -> None:
        celulas = tabela.add_row().cells
        celulas[0].text = rotulo
        celulas[1].text = valor

    _linha("Data", reuniao.criado_em.strftime("%d/%m/%Y"))
    if reuniao.duracao_segundos:
        _linha("Duração", _formatar_tempo(reuniao.duracao_segundos))
    participantes = sorted({s.falante for s in reuniao.segmentos if s.falante})
    _linha("Participantes", ", ".join(participantes) if participantes else "Não identificados individualmente")

    documento.add_paragraph()
    _adicionar_secao(documento, "Objetivo", ata.objetivo, "Não preenchido.")
    _adicionar_secao(documento, "Assuntos tratados", ata.assuntos_tratados, "Não preenchido.")
    _adicionar_secao(documento, "Decisões", ata.decisoes, "Nenhuma decisão formal registrada.")

    documento.add_heading("Encaminhamentos", level=2)
    if not encaminhamentos:
        documento.add_paragraph("Nenhum encaminhamento registrado.")
    else:
        tabela_enc = documento.add_table(rows=1, cols=4)
        tabela_enc.style = "Light Grid Accent 1"
        cabecalho = tabela_enc.rows[0].cells
        cabecalho[0].text = "Encaminhamento"
        cabecalho[1].text = "Responsável"
        cabecalho[2].text = "Prazo"
        cabecalho[3].text = "Status"
        for enc in encaminhamentos:
            celulas = tabela_enc.add_row().cells
            celulas[0].text = enc.descricao
            celulas[1].text = enc.responsavel or "Não identificado"
            celulas[2].text = enc.prazo or "Não identificado"
            celulas[3].text = enc.status or "Pendente"

    documento.add_paragraph()
    _adicionar_secao(documento, "Pendências", ata.pendencias, "Nenhuma pendência registrada.")

    rodape = documento.add_paragraph()
    rodape.add_run(
        "\nDocumento gerado com apoio do Transcritor CTCE a partir de processamento "
        "local de áudio. Recomenda-se revisão humana antes do uso oficial deste conteúdo."
    ).italic = True

    buffer = BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer
